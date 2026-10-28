#!/usr/bin/env python3
"""Build a versioned CV design example. All writes stay beside this script.

No Word, Google Docs, browser, network access, installation, or review agent is
used. The PDF is laid out by ReportLab; the DOCX is independently authored with
real paragraph styles and numbering. Its pagination needs a Word/Docs check.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import importlib.metadata
import io
import json
import re
import sys
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime, timezone
from pathlib import Path

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parent
VENDORED_PACKAGES = ROOT.parent / "tools" / "python"
if VENDORED_PACKAGES.is_dir():
    sys.path.insert(0, str(VENDORED_PACKAGES))

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fontTools.ttLib import TTFont as InspectFont
from lxml import etree
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph


PREFIX = "round-1"
PAGE_W, PAGE_H = A4
MARGIN = 18 * mm
LIVE_W = PAGE_W - 2 * MARGIN
FOOTER_RESERVE = 34.0
FRAME_BOTTOM = MARGIN + FOOTER_RESERVE
FRAME_TOP = PAGE_H - MARGIN
FOOTER_BASELINE = MARGIN + 3.5
FOOTER_RULE_Y = MARGIN + 22.0
BULLET_INDENT = 4 * mm
PNG_DPI = 144
INK = "202725"
MUTED = "515A57"
BLUE = "294FA8"
RULE = "D7DDD9"
FOOTER_TEXT = "WORKING DRAFT · Not for submission"
TITLE = "Kwanho Kim — CV design example — Working draft"
FROZEN_TIME = datetime(2026, 10, 26, tzinfo=timezone.utc)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@dataclass(frozen=True)
class Spec:
    size: float = 10.5
    leading: float = 14.5
    before: float = 0
    after: float = 0
    color: str = INK
    bold: bool = False
    italic: bool = False
    keep_next: bool = False
    left: float = 0
    first_line: float = 0


# Points, except colors. Adjacent paragraph spaces overlap (max, not sum).
STYLES = {
    "Normal": Spec(),
    "Title": Spec(size=31, leading=34, after=1, bold=True, keep_next=True),
    "Subtitle": Spec(size=12, leading=16, after=5, color=BLUE, keep_next=True),
    "CV Contact": Spec(size=9.5, leading=13, color=BLUE, keep_next=True),
    "CV Draft": Spec(
        size=9.25, leading=12, before=7, after=1, color=BLUE,
        bold=True, keep_next=True,
    ),
    "CV Note": Spec(size=9.25, leading=12, color=MUTED, italic=True),
    "Heading 1": Spec(
        size=10.5, leading=14.5, before=12, after=5,
        color=BLUE, bold=True, keep_next=True,
    ),
    "Heading 2": Spec(size=11.25, leading=14.5, before=8, after=1, keep_next=True),
    "CV Metadata": Spec(
        size=9.25, leading=12.5, after=3, color=MUTED, keep_next=True,
    ),
    "CV Bullet": Spec(after=2.5, left=BULLET_INDENT, first_line=-BULLET_INDENT),
    "CV Skill": Spec(after=3),
    "CV Continuation": Spec(size=12, leading=16, after=2, keep_next=True),
    "CV Record": Spec(before=6, after=1, keep_next=True),
    "CV Record Metadata": Spec(size=9.25, leading=12.5, color=MUTED),
    "CV Footer": Spec(size=9.25, leading=12, color=BLUE),
}


@dataclass
class Block:
    id: str
    style: str
    runs: list[dict] = field(default_factory=list)
    bullet: bool = False
    space_before: float | None = None

    @property
    def text(self) -> str:
        return "".join(run["text"] for run in self.runs)


def text_block(key: str, style: str, text: str, **kwargs) -> Block:
    return Block(key, style, [{"text": text}], **kwargs)


def paired_block(key: str, style: str, first: str, second: str) -> Block:
    return Block(
        key, style,
        [{"text": first, "bold": True}, {"text": " — " + second}],
    )


def make_pages(data: dict) -> list[list[Block]]:
    """The single reading sequence consumed by all four document formats."""
    person = data["identity"]
    p1 = [
        text_block("name", "Title", person["name"]),
        text_block("headline", "Subtitle", person["headline"]),
    ]
    for index, row in enumerate(person["contacts"], start=1):
        runs = []
        for contact in row:
            if runs:
                runs.append({"text": " · "})
            runs.append(contact.copy())
        p1.append(Block(f"contacts-{index}", "CV Contact", runs))
    if data.get("notice"):
        p1.extend([
            text_block("draft-label", "CV Draft", data["notice"]["label"]),
            text_block("draft-note", "CV Note", data["notice"]["note"]),
        ])
    p1.extend([
        text_block("summary-heading", "Heading 1", "SUMMARY"),
        text_block("summary", "Normal", data["summary"]),
        text_block("experience-heading", "Heading 1", "EXPERIENCE"),
    ])
    for index, job in enumerate(data["experience"], start=1):
        key = f"job-{index}"
        p1.extend([
            paired_block(key, "Heading 2", job["employer"], job["role"]),
            text_block(f"{key}-metadata", "CV Metadata", job["metadata"]),
        ])
        p1.extend(
            text_block(f"{key}-bullet-{n}", "CV Bullet", text, bullet=True)
            for n, text in enumerate(job["bullets"], start=1)
        )
    p1.append(text_block("skills-heading", "Heading 1", "SKILLS"))
    for index, skill in enumerate(data["skills"], start=1):
        p1.append(Block(f"skill-{index}", "CV Skill", [
            {"text": skill["label"] + ": ", "bold": True},
            {"text": skill["text"]},
        ]))

    p2 = [
        paired_block("continuation", "CV Continuation", person["name"], person["headline"]),
        text_block("projects-heading", "Heading 1", "SELECTED PROJECTS"),
    ]
    for index, project in enumerate(data["projects"], start=1):
        key = f"project-{index}"
        p2.append(paired_block(key, "Heading 2", project["name"], project["descriptor"]))
        metadata = [{"text": project["metadata"]}]
        if project.get("link"):
            metadata.extend([{"text": " · "}, project["link"].copy()])
        p2.append(Block(f"{key}-metadata", "CV Metadata", metadata))
        p2.extend(
            text_block(f"{key}-bullet-{n}", "CV Bullet", text, bullet=True)
            for n, text in enumerate(project["bullets"], start=1)
        )
    p2.append(text_block("education-heading", "Heading 1", "EDUCATION"))
    if data["education"].get("note"):
        p2.append(text_block("education-note", "CV Note", data["education"]["note"]))
    for index, entry in enumerate(data["education"]["entries"], start=1):
        key = f"education-{index}"
        p2.extend([
            paired_block(key, "CV Record", entry["subject"], entry["institution"]),
            text_block(f"{key}-metadata", "CV Record Metadata", entry["dates"]),
        ])
    training_title = "TRAINING & CERTIFICATION" if PREFIX != "round-1" else "TRAINING"
    p2.append(text_block("training-heading", "Heading 1", training_title))
    for index, entry in enumerate(data["training"], start=1):
        key = f"training-{index}"
        p2.extend([
            paired_block(key, "CV Record", entry["name"], entry["description"]),
            text_block(f"{key}-metadata", "CV Record Metadata", entry["metadata"]),
        ])
    if PREFIX == "round-1":
        p2.append(text_block("credential-heading", "Heading 1", "SELECTED CREDENTIAL"))
    p2.extend([
        text_block("credential", "CV Record", data["credential"]["name"]),
        text_block("credential-metadata", "CV Record Metadata", data["credential"]["metadata"]),
        text_block("leadership-heading", "Heading 1", "LEADERSHIP"),
        paired_block("leadership", "CV Record", data["leadership"]["name"], data["leadership"]["role"]),
        text_block("leadership-metadata", "CV Metadata", data["leadership"]["metadata"]),
        text_block("leadership-bullet", "CV Bullet", data["leadership"]["bullet"], bullet=True),
    ])
    # A single-name credential is a record label, not an unweighted body line.
    next(b for b in p2 if b.id == "credential").runs[0]["bold"] = True
    if PREFIX != "round-1":
        next(b for b in p1 if b.id == "job-1").space_before = 6
        next(b for b in p2 if b.id == "project-1").space_before = 6
    return [p1, p2]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font_info(pages: list[list[Block]]) -> list[dict]:
    wanted = set("".join(b.text for p in pages for b in p) + FOOTER_TEXT + "•12 /")
    wanted = {ord(c) for c in wanted if not c.isspace()}
    results = []
    for name in ("Regular", "Bold", "Italic"):
        path = ROOT / "fonts" / f"Lato-{name}.ttf"
        with InspectFont(path) as font:
            missing = sorted(wanted - set(font.getBestCmap()))
            results.append({
                "file": str(path.relative_to(ROOT)),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
                "embedding_fsType": font["OS/2"].fsType,
                "missing_required_codepoints": [f"U+{cp:04X}" for cp in missing],
            })
    return results


def register_pdf_fonts() -> None:
    for logical_name, filename in [
        ("Lato", "Lato-Regular.ttf"),
        ("Lato-Bold", "Lato-Bold.ttf"),
        ("Lato-Italic", "Lato-Italic.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(logical_name, str(ROOT / "fonts" / filename)))
    pdfmetrics.registerFontFamily(
        "Lato", normal="Lato", bold="Lato-Bold", italic="Lato-Italic",
        boldItalic="Lato-Bold",  # No mixed bold/italic run is used.
    )


def pdf_styles() -> dict[str, ParagraphStyle]:
    result = {}
    for name, spec in STYLES.items():
        font = "Lato-Bold" if spec.bold else "Lato-Italic" if spec.italic else "Lato"
        result[name] = ParagraphStyle(
            name,
            fontName=font,
            fontSize=spec.size,
            leading=spec.leading,
            textColor=colors.HexColor("#" + spec.color),
            alignment=TA_LEFT,
            spaceBefore=spec.before,
            spaceAfter=spec.after,
            leftIndent=spec.left,
            # ReportLab labels occupy their own bullet box. Continuation lines
            # align at leftIndent; do not insert spaces or a bullet into text.
            firstLineIndent=0,
            bulletIndent=0,
            bulletFontName="Lato",
            # A 9.5 pt bullet fits the 4 mm label box. Using a 10.5 pt label
            # triggers ReportLab's minimum bullet gap and shifts line one.
            bulletFontSize=9.5,
            keepWithNext=spec.keep_next,
            allowWidows=0,
            allowOrphans=0,
            splitLongWords=0,
            uriWasteReduce=0,
        )
    return result


def pdf_markup(block: Block) -> str:
    fragments = []
    for run in block.runs:
        text = html.escape(run["text"])
        if run.get("bold"):
            text = f"<b>{text}</b>"
        if run.get("italic"):
            text = f"<i>{text}</i>"
        if run.get("href"):
            uri = html.escape(run["href"], quote=True)
            text = f'<link href="{uri}" color="#{BLUE}"><u>{text}</u></link>'
        fragments.append(text)
    return "".join(fragments)


class TrackedParagraph(Paragraph):
    """Record flowable positions as well as independently inspected glyph boxes."""

    def __init__(self, block: Block, style: ParagraphStyle, placements: list[dict]):
        if block.space_before is not None:
            style = ParagraphStyle(
                f"{style.name}-{block.id}", parent=style, spaceBefore=block.space_before,
            )
        super().__init__(pdf_markup(block), style, bulletText="•" if block.bullet else None)
        self.block = block
        self.placements = placements

    def split(self, availWidth, availHeight):
        # Keep individual one-to-three-line paragraphs whole. The document
        # still permits splitting KeepTogether wrappers used by keepWithNext.
        return []

    def drawOn(self, canvas, x, y, _sW=0):
        self.placements.append({
            "id": self.block.id,
            "page": canvas.getPageNumber(),
            "style": self.block.style,
            "space_before_pt": self.style.spaceBefore,
            "paragraph_bbox_pt_top_origin": [
                round(x, 3), round(PAGE_H - y - self.height, 3),
                round(x + self.width, 3), round(PAGE_H - y, 3),
            ],
            "line_count": len(self.blPara.lines),
            "text": self.block.text,
        })
        return super().drawOn(canvas, x, y, _sW)


def draw_footer(canvas: Canvas, document: BaseDocTemplate) -> None:
    # onPageEnd is intentional: the text stream must be body, then footer.
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#" + RULE))
    canvas.setLineWidth(0.5)
    canvas.line(MARGIN, FOOTER_RULE_Y, PAGE_W - MARGIN, FOOTER_RULE_Y)
    canvas.setFont("Lato", 9.25)
    canvas.setFillColor(colors.HexColor("#" + STYLES["CV Footer"].color))
    canvas.drawString(MARGIN, FOOTER_BASELINE, FOOTER_TEXT)
    canvas.setFillColor(colors.HexColor("#" + MUTED))
    canvas.drawRightString(PAGE_W - MARGIN, FOOTER_BASELINE, f"{document.page} / 2")
    canvas.restoreState()


def build_pdf(pages: list[list[Block]]) -> list[dict]:
    register_pdf_fonts()
    styles = pdf_styles()
    placements: list[dict] = []
    doc = BaseDocTemplate(
        str(ROOT / f"{PREFIX}.pdf"),
        pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN, bottomMargin=MARGIN,
        allowSplitting=1,
        title=TITLE,
        author="Kwanho Kim",
        subject="C++ software engineering. Career facts supplied by the owner." if PREFIX == "current" else "CV design example. Role and education details await confirmation.",
        creator="build_cv.py · ReportLab",
        keywords="CV, Software Engineer, Working draft, Design example",
        invariant=1,
        pageCompression=1,
        initialFontName="Lato",
        initialFontSize=10.5,
        initialLeading=14.5,
        lang="en-GB",
    )
    frame = Frame(
        MARGIN, FRAME_BOTTOM, LIVE_W, FRAME_TOP - FRAME_BOTTOM,
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
        id="single-reading-flow", overlapAttachedSpace=1,
    )
    doc.addPageTemplates(PageTemplate(id="A4", frames=[frame], onPageEnd=draw_footer))
    story = []
    for index, blocks in enumerate(pages):
        if index:
            story.append(PageBreak())
        story.extend(TrackedParagraph(block, styles[block.style], placements) for block in blocks)
    doc.build(story, canvasmaker=Canvas)
    return placements


def set_lato(font, size: float | None = None) -> None:
    font.name = "Lato"
    if size is not None:
        font.size = Pt(size)
    rpr = font._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), "Lato")
    # Do not leave a theme font overriding the explicitly selected family.
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)


def configure_docx_styles(document) -> None:
    for name, spec in STYLES.items():
        if name in document.styles:
            style = document.styles[name]
        else:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != "Normal":
            style.base_style = document.styles["Normal"]
        set_lato(style.font, spec.size)
        style.font.bold = spec.bold
        style.font.italic = spec.italic
        style.font.color.rgb = RGBColor.from_string(spec.color)
        pf = style.paragraph_format
        pf.space_before = Pt(spec.before)
        pf.space_after = Pt(spec.after)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(spec.leading)
        pf.left_indent = Pt(spec.left)
        pf.right_indent = Pt(0)
        pf.first_line_indent = Pt(spec.first_line)
        pf.keep_with_next = spec.keep_next
        pf.keep_together = True
        pf.widow_control = True
        pf.page_break_before = False
        # Explicit language for spellchecking, not a claim of formal tagging.
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "en-GB")
        style.element.get_or_add_rPr().append(lang)
    if "Hyperlink" not in document.styles:
        document.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    link_style = document.styles["Hyperlink"]
    set_lato(link_style.font)
    link_style.font.color.rgb = RGBColor.from_string(BLUE)
    link_style.font.underline = True


def add_bullet_definition(document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in [("start", "1"), ("numFmt", "bullet"), ("lvlText", "•"), ("lvlJc", "left"), ("suff", "tab")]:
        child = OxmlElement("w:" + tag)
        child.set(qn("w:val"), value)
        lvl.append(child)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(round(BULLET_INDENT * 20)))
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(round(BULLET_INDENT * 20)))
    ind.set(qn("w:hanging"), str(round(BULLET_INDENT * 20)))
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi"):
        fonts.set(qn("w:" + key), "Lato")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    rpr.append(size)
    lvl.append(rpr)
    abstract.append(lvl)
    # abstractNum elements precede num elements in the OOXML sequence.
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def docx_link(paragraph, run_data: dict) -> None:
    relationship = paragraph.part.relate_to(
        run_data["href"],
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = run_data["text"]
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def docx_field(paragraph, instruction: str, cached: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(begin)
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    paragraph.add_run()._r.append(code)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(cached)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def deterministic_docx_save(document, path: Path) -> None:
    """Normalize ZIP timestamps; no scratch directory or external writes."""
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            output = zipfile.ZipInfo(item.filename, date_time=(2026, 10, 26, 0, 0, 0))
            output.compress_type = zipfile.ZIP_DEFLATED
            output.external_attr = item.external_attr
            target.writestr(output, source.read(item.filename))


def build_docx(pages: list[list[Block]]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    section.header_distance = Mm(18)
    section.footer_distance = Mm(18)
    configure_docx_styles(document)
    bullet_id = add_bullet_definition(document)
    for page_index, blocks in enumerate(pages):
        for block_index, block in enumerate(blocks):
            p = document.add_paragraph(style=block.style)
            if block.space_before is not None:
                p.paragraph_format.space_before = Pt(block.space_before)
            if page_index and block_index == 0:
                p.paragraph_format.page_break_before = True
            if block.bullet:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = 0
                numpr.get_or_add_numId().val = bullet_id
            for run_data in block.runs:
                if run_data.get("href"):
                    docx_link(p, run_data)
                    continue
                run = p.add_run(run_data["text"])
                if run_data.get("bold"):
                    run.bold = True
                if run_data.get("italic"):
                    run.italic = True
    footer = section.footer.paragraphs[0]
    footer.style = document.styles["CV Footer"]
    footer.paragraph_format.tab_stops.add_tab_stop(Pt(LIVE_W), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run(FOOTER_TEXT + "\t")
    docx_field(footer, "PAGE", "1")
    footer.add_run(" / ")
    docx_field(footer, "NUMPAGES", "2")
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    for key, val in [("val", "single"), ("sz", "4"), ("space", "8"), ("color", RULE)]:
        top.set(qn("w:" + key), val)
    borders.append(top)
    footer._p.get_or_add_pPr().append(borders)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    core = document.core_properties
    core.title = TITLE
    core.subject = "C++ software engineering; owner-confirmed career facts." if PREFIX == "current" else "CV design example; role and education details await confirmation."
    core.author = "Kwanho Kim"
    core.keywords = "CV, C++, Software Engineer" if PREFIX == "current" else "CV, Software Engineer, Working draft, Design example"
    core.comments = "Career facts updated from the owner's answers on 26 October 2026." if PREFIX == "current" else "Design example; see the accompanying review record."
    core.created = core.modified = FROZEN_TIME
    deterministic_docx_save(document, ROOT / f"{PREFIX}.docx")


def markdown_runs(block: Block) -> str:
    parts = []
    for run in block.runs:
        text = run["text"]
        if run.get("href"):
            text = f'[{text}]({run["href"]})'
        elif run.get("bold"):
            text = re.sub(r"^(\s*)(.*?)(\s*)$", r"\1**\2**\3", text)
        parts.append(text)
    return "".join(parts)


def build_text_copies(pages: list[list[Block]]) -> None:
    md, plain = [], []
    for index, blocks in enumerate(pages):
        if index:
            md.extend(["<!-- Deliberate page break: page 2 -->", ""])
            plain.extend(["--- PLANNED PAGE BREAK ---", ""])
        for block in blocks:
            rich = markdown_runs(block)
            if block.style == "Title":
                md.append("# " + rich)
            elif block.style == "Heading 1":
                md.append("## " + rich)
            elif block.style == "Heading 2":
                md.append("### " + rich)
            elif block.style in ("CV Draft", "CV Note"):
                md.append("> " + rich)
            elif block.bullet:
                md.append("- " + rich)
            else:
                md.append(rich)
            md.append("")
            plain.extend([("- " if block.bullet else "") + block.text, ""])
        footer = f"{FOOTER_TEXT} | {index + 1} / 2"
        md.extend(["_" + footer + "_", ""])
        plain.extend([footer, ""])
    (ROOT / f"{PREFIX}.md").write_text("\n".join(md), encoding="utf-8")
    (ROOT / f"{PREFIX}.txt").write_text("\n".join(plain), encoding="utf-8")


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def within_live_area(box, tolerance: float = 0.05) -> bool:
    x0, y0, x1, y1 = box
    return (
        x0 >= MARGIN - tolerance and y0 >= MARGIN - tolerance
        and x1 <= PAGE_W - MARGIN + tolerance
        and y1 <= PAGE_H - MARGIN + tolerance
    )


def check_pdf(pages: list[list[Block]], placements: list[dict]) -> dict:
    document = pymupdf.open(ROOT / f"{PREFIX}.pdf")
    page_reports, font_xrefs, found_links, errors = [], set(), [], []
    extraction = []
    for index, page in enumerate(document):
        text = page.get_text("text", sort=False)
        extraction.append(f"--- PAGE {index + 1} ---\n{text}")
        spans = [
            span for block in page.get_text("dict", sort=False)["blocks"]
            if block["type"] == 0 for line in block["lines"] for span in line["spans"]
            if span["text"].strip()
        ]
        bad_boxes = [
            {"text": span["text"], "bbox_pt": list(span["bbox"])}
            for span in spans if not within_live_area(span["bbox"])
        ]
        bad_drawing_boxes = [
            list(drawing["rect"]) for drawing in page.get_drawings()
            if not within_live_area(drawing["rect"])
        ]
        cursor, matches, order_ok = 0, [], True
        actual = normalized(text)
        for block in pages[index] if index < len(pages) else []:
            needle = normalized(block.text)
            start = actual.find(needle, cursor)
            matches.append({"id": block.id, "found_at": start})
            if start < 0:
                order_ok = False
                errors.append(f"Page {index + 1}: missing/out-of-order text for {block.id}")
            else:
                cursor = start + len(needle)
        footer_after_body = actual.find(FOOTER_TEXT, cursor) >= cursor
        content_spans = [s for s in spans if s["bbox"][1] < PAGE_H - FOOTER_RULE_Y]
        content_bottom = max((s["bbox"][3] for s in content_spans), default=0)
        all_boxes = [s["bbox"] for s in spans]
        union = [
            min((b[0] for b in all_boxes), default=0),
            min((b[1] for b in all_boxes), default=0),
            max((b[2] for b in all_boxes), default=0),
            max((b[3] for b in all_boxes), default=0),
        ]
        for font in page.get_fonts(full=True):
            font_xrefs.add(font[0])
        links = [
            {"uri": link.get("uri"), "bbox_pt": list(link["from"]),
             "within_margins": within_live_area(link["from"])}
            for link in page.get_links()
        ]
        found_links.extend(link["uri"] for link in links if link["uri"])
        page_reports.append({
            "page": index + 1,
            "size_pt": [round(page.rect.width, 3), round(page.rect.height, 3)],
            "a4": abs(page.rect.width - PAGE_W) < 0.1 and abs(page.rect.height - PAGE_H) < 0.1,
            "text_span_count": len(spans),
            "glyph_bbox_union_pt": [round(v, 3) for v in union],
            "out_of_margin_spans": bad_boxes,
            "out_of_margin_drawing_boxes": bad_drawing_boxes,
            "content_ink_bottom_pt": round(content_bottom, 3),
            "gap_from_content_ink_to_footer_rule_pt": round(PAGE_H - FOOTER_RULE_Y - content_bottom, 3),
            "reading_order_passed": order_ok and bool(matches),
            "block_matches": matches,
            "footer_after_body": footer_after_body,
            "footer_colour_matches_style": any(
                s["text"] == FOOTER_TEXT and s["color"] == int(STYLES["CV Footer"].color, 16)
                for s in spans
            ),
            "replacement_character_absent": "\ufffd" not in text,
            "links": links,
        })
        pixmap = page.get_pixmap(dpi=PNG_DPI, alpha=False)
        png_path = ROOT / f"{PREFIX}-page-{index + 1}.png"
        pixmap.save(png_path)
    font_reports = []
    for xref in sorted(font_xrefs):
        name, extension, kind, binary = document.extract_font(xref)
        unicode_type, unicode_value = document.xref_get_key(xref, "ToUnicode")
        font_reports.append({
            "xref": xref, "name": name, "extension": extension, "type": kind,
            "embedded_bytes": len(binary),
            "embedded_sha256": hashlib.sha256(binary).hexdigest() if binary else None,
            "to_unicode": {"type": unicode_type, "value": unicode_value},
        })
    expected_links = sorted(
        run["href"] for blocks in pages for block in blocks for run in block.runs
        if run.get("href")
    )
    placement_order = [
        (page_index + 1, block.id)
        for page_index, blocks in enumerate(pages) for block in blocks
    ] == [(item["page"], item["id"]) for item in placements]
    checks = {
        "two_pages": len(document) == 2,
        "all_pages_a4": all(p["a4"] for p in page_reports),
        "all_text_within_18mm_margins": all(not p["out_of_margin_spans"] for p in page_reports),
        "rules_and_underlines_within_margins": all(not p["out_of_margin_drawing_boxes"] for p in page_reports),
        "all_expected_text_in_linear_order": all(p["reading_order_passed"] for p in page_reports),
        "flowable_order_matches_planned_pages": placement_order,
        "footers_follow_body_in_extraction": all(p["footer_after_body"] for p in page_reports),
        "footer_colour_matches_style": all(p["footer_colour_matches_style"] for p in page_reports),
        "no_unicode_replacement_character": all(p["replacement_character_absent"] for p in page_reports),
        "nonempty_embedded_font_programs": bool(font_reports) and all(f["embedded_bytes"] > 0 for f in font_reports),
        "all_fonts_have_unicode_maps": all(f["to_unicode"]["type"] == "xref" for f in font_reports),
        "expected_hyperlink_annotations": sorted(found_links) == expected_links,
        "link_rectangles_within_margins": all(l["within_margins"] for p in page_reports for l in p["links"]),
        "title_metadata": document.metadata.get("title") == TITLE,
    }
    metadata = document.metadata.copy()
    document.close()
    (ROOT / f"{PREFIX}-extracted.txt").write_text("\n".join(extraction), encoding="utf-8")
    return {
        "checks": checks, "page_count": len(page_reports),
        "pages": page_reports, "fonts": font_reports,
        "expected_link_uris": expected_links, "metadata": metadata,
        "placements": placements, "errors": errors,
    }


def check_docx(pages: list[list[Block]]) -> dict:
    with zipfile.ZipFile(ROOT / f"{PREFIX}.docx") as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))
        relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
        footer = etree.fromstring(archive.read("word/footer1.xml"))
        settings = etree.fromstring(archive.read("word/settings.xml"))
    paragraphs = document.xpath("/w:document/w:body/w:p", namespaces=NS)
    body_text = ["".join(p.xpath(".//w:t/text()", namespaces=NS)) for p in paragraphs]
    expected_text = [block.text for blocks in pages for block in blocks]
    page_size = document.xpath("//w:sectPr/w:pgSz", namespaces=NS)[0]
    page_margins = document.xpath("//w:sectPr/w:pgMar", namespaces=NS)[0]
    link_targets = sorted(
        r.get("Target") for r in relationships
        if r.get("Type", "").endswith("/hyperlink") and r.get("TargetMode") == "External"
    )
    expected_links = sorted(
        run["href"] for blocks in pages for block in blocks for run in block.runs
        if run.get("href")
    )
    page_breaks = [
        element for element in document.xpath("//w:pPr/w:pageBreakBefore", namespaces=NS)
        if element.get(qn("w:val"), "1") not in ("0", "false", "off")
    ]
    bullets = document.xpath("//w:pPr/w:numPr", namespaces=NS)
    style_results = {}
    for name in STYLES:
        matches = styles.xpath(
            "//w:style[w:name/@w:val=$name]", namespaces=NS,
            name={"Heading 1": "heading 1", "Heading 2": "heading 2"}.get(name, name),
        )
        if not matches:
            style_results[name] = {"present": False}
            continue
        element = matches[0]
        spacing = element.find("w:pPr/w:spacing", NS)
        widows = element.find("w:pPr/w:widowControl", NS)
        keep_next = element.find("w:pPr/w:keepNext", NS)
        font = element.find("w:rPr/w:rFonts", NS)
        style_results[name] = {
            "present": True,
            "font": font.get(qn("w:ascii")) if font is not None else None,
            "line_spacing_twips": spacing.get(qn("w:line")) if spacing is not None else None,
            "line_spacing_rule": spacing.get(qn("w:lineRule")) if spacing is not None else None,
            "widow_control": widows is not None and widows.get(qn("w:val"), "1") != "0",
            "keep_with_next": keep_next is not None and keep_next.get(qn("w:val"), "1") != "0",
        }
    hanging = numbering.xpath(
        "//w:lvl[w:numFmt/@w:val='bullet'][w:lvlText/@w:val='•']/w:pPr/w:ind",
        namespaces=NS,
    )
    bullet_indent_ok = any(
        e.get(qn("w:left")) == str(round(BULLET_INDENT * 20))
        and e.get(qn("w:hanging")) == str(round(BULLET_INDENT * 20))
        for e in hanging
    )
    checks = {
        "a4_geometry": page_size.get(qn("w:w")) == str(round(210 * mm * 20))
        and page_size.get(qn("w:h")) == str(round(297 * mm * 20)),
        "18mm_margins": all(
            page_margins.get(qn("w:" + side)) == str(round(MARGIN * 20))
            for side in ("top", "right", "bottom", "left")
        ),
        "one_deliberate_page_break": len(page_breaks) == 1,
        "editable_paragraph_text_matches_copy": body_text == expected_text,
        "all_named_styles_present": all(v["present"] for v in style_results.values()),
        "lato_in_all_styles": all(v.get("font") == "Lato" for v in style_results.values()),
        "exact_line_spacing_in_all_styles": all(v.get("line_spacing_rule") == "exact" for v in style_results.values()),
        "widow_control_in_all_styles": all(v.get("widow_control") for v in style_results.values()),
        "headings_keep_with_next": all(
            style_results[s].get("keep_with_next") for s in ("Heading 1", "Heading 2")
        ),
        "real_numbered_bullet_paragraphs": len(bullets) == sum(b.bullet for p in pages for b in p),
        "real_4mm_hanging_bullets": bullet_indent_ok,
        "no_tables_or_text_boxes_or_drawings": not document.xpath(
            "//w:tbl | //w:txbxContent | //w:drawing | //w:pict", namespaces=NS,
        ),
        "expected_external_hyperlinks": link_targets == expected_links,
        "footer_text": FOOTER_TEXT in "".join(footer.xpath("//w:t/text()", namespaces=NS)),
        "page_number_fields": {
            s.strip() for s in footer.xpath("//w:instrText/text()", namespaces=NS)
        } == {"PAGE", "NUMPAGES"},
        "update_fields_requested": bool(settings.xpath("//w:updateFields[@w:val='true']", namespaces=NS)),
    }
    return {
        "checks": checks,
        "body_paragraph_count": len(paragraphs),
        "bullet_paragraph_count": len(bullets),
        "explicit_page_break_count": len(page_breaks),
        "external_hyperlink_count": len(link_targets),
        "styles": style_results,
        "rendered_page_count": None,
        "pagination_status": "Not rendered in Word or Google Docs; two pages are planned, not independently verified.",
        "font_embedding": "DOCX references Lato; font files are supplied separately, not embedded in this DOCX.",
    }


def contrast_against_white(hex_color: str) -> float:
    components = [int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    linear = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in components]
    luminance = sum(v * weight for v, weight in zip(linear, (0.2126, 0.7152, 0.0722)))
    return round(1.05 / (luminance + 0.05), 2)


def main() -> int:
    global PREFIX, FOOTER_TEXT, TITLE, BLUE
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--version", choices=("round-1", "round-2", "current"), default="current")
    parser.add_argument(
        "--check-repeatability", action="store_true",
        help="Build the selected version twice and compare its SHA-256 hashes.",
    )
    args = parser.parse_args()
    PREFIX = args.version
    if PREFIX == "current":
        old_accent = BLUE
        BLUE = "20543E"
        FOOTER_TEXT = "Kwanho Kim · C++ Software Engineer"
        TITLE = "Kwanho Kim — CV"
        for name, style in list(STYLES.items()):
            if style.color == old_accent:
                STYLES[name] = replace(style, color=BLUE)
    if PREFIX != "round-1":
        STYLES.update({
            "CV Draft": replace(STYLES["CV Draft"], color=MUTED),
            "CV Footer": replace(STYLES["CV Footer"], color=MUTED),
            "Heading 1": replace(STYLES["Heading 1"], size=11, leading=15, before=14),
            "Heading 2": replace(STYLES["Heading 2"], size=12, leading=16, before=12, after=2),
            "CV Metadata": replace(STYLES["CV Metadata"], size=9.5),
            "CV Record Metadata": replace(STYLES["CV Record Metadata"], size=9.5),
            "CV Note": replace(STYLES["CV Note"], size=9.5, leading=12.5),
        })
    data = json.loads((ROOT / f"{PREFIX}-content.json").read_text(encoding="utf-8"))
    if data["version"] != PREFIX:
        raise ValueError("The content version must match the requested output version.")
    pages = make_pages(data)
    source_fonts = font_info(pages)
    if any(font["missing_required_codepoints"] for font in source_fonts):
        raise ValueError("A supplied Lato font is missing required Unicode characters.")
    placements = build_pdf(pages)
    build_docx(pages)
    build_text_copies(pages)
    pdf_report = check_pdf(pages, placements)
    docx_report = check_docx(pages)
    artifact_names = [
        f"{PREFIX}.pdf", f"{PREFIX}.docx", f"{PREFIX}.md", f"{PREFIX}.txt",
        f"{PREFIX}-extracted.txt", f"{PREFIX}-page-1.png", f"{PREFIX}-page-2.png",
    ]
    repeatability = {
        "performed": False,
        "note": "Use --check-repeatability for an actual two-build byte comparison.",
    }
    if args.check_repeatability:
        original_hashes = {name: sha256(ROOT / name) for name in artifact_names}
        placements = build_pdf(pages)
        build_docx(pages)
        build_text_copies(pages)
        pdf_report = check_pdf(pages, placements)
        docx_report = check_docx(pages)
        repeatability = {
            "performed": True,
            "method": "Two consecutive local builds; SHA-256 comparison of each generated artifact.",
            "artifact_hash_matches": {
                name: original_hashes[name] == sha256(ROOT / name) for name in artifact_names
            },
            "note": "Verification JSON intentionally includes the current build time and is not compared.",
        }
    artifacts = []
    for name in artifact_names:
        path = ROOT / name
        entry = {"file": name, "bytes": path.stat().st_size, "sha256": sha256(path)}
        if path.suffix == ".png":
            with Image.open(path) as image:
                entry["pixel_size"] = list(image.size)
            entry["render_dpi"] = PNG_DPI
        artifacts.append(entry)
    passed = all(pdf_report["checks"].values()) and all(docx_report["checks"].values())
    if args.check_repeatability:
        passed = passed and all(repeatability["artifact_hash_matches"].values())
    report = {
        "version": PREFIX,
        "purpose": "CV using owner-confirmed facts and the reviewed reference layout." if PREFIX == "current" else "CV design example; build checks are separate from design review.",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "automated_build_checks_passed": passed,
        "owner_facts_confirmed": PREFIX == "current",
        "geometry": {
            "page_mm": [210, 297], "margin_mm": 18,
            "text_width_mm": 174, "body_font_pt": 10.5, "body_leading_pt": 14.5,
            "metadata_font_pt": STYLES["CV Metadata"].size, "bullet_indent_mm": 4,
            "body_frame_top_pt": round(MARGIN, 3),
            "body_frame_bottom_pt_top_origin": round(PAGE_H - FRAME_BOTTOM, 3),
            "footer_rule_pt_top_origin": round(PAGE_H - FOOTER_RULE_Y, 3),
            "bbox_tolerance_pt": 0.05,
        },
        "source_content_sha256": sha256(ROOT / f"{PREFIX}-content.json"),
        "build_script_sha256": sha256(Path(__file__).resolve()),
        "style_specifications_pt": {name: asdict(spec) for name, spec in STYLES.items()},
        "environment": {
            "python": sys.version.split()[0],
            **{name: importlib.metadata.version(name) for name in (
                "reportlab", "python-docx", "pymupdf", "Pillow", "fonttools",
            )},
        },
        "source_fonts": source_fonts,
        "text_contrast_against_white": {
            "#" + value: contrast_against_white(value) for value in (INK, MUTED, BLUE)
        },
        "pdf": pdf_report,
        "docx": docx_report,
        "artifacts": artifacts,
        "repeatability": repeatability,
        "limitations": [
            "No authenticated Google Docs session, font-menu check, import or export test.",
            "No Word/LibreOffice rendering: DOCX page count and visual parity are not verified.",
            "No real ATS ingestion or recruiter-system test; extraction is a local PyMuPDF check only.",
            "The PDF contains selectable Unicode text but is not a tagged PDF/UA document.",
            "Hyperlink annotation targets are checked; remote availability is not tested.",
            "Facts supplied by the owner; no independent employer or university verification." if PREFIX == "current" else "Role details, personal contribution, education, award scope and project status remain unresolved.",
            "No formal design evaluation round has been performed by this build.",
        ],
    }
    output = ROOT / f"{PREFIX}-verification.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "automated_build_checks_passed": passed,
        "pdf_page_count": pdf_report["page_count"],
        "pdf_checks": pdf_report["checks"],
        "docx_checks": docx_report["checks"],
        "repeatability": repeatability,
        "verification": str(output),
    }, ensure_ascii=False, indent=2))
    return 0 if passed else 1


if __name__ == "__main__":
    try:
        exit_code = main()
    except Exception as exc:
        # Never leave a previous successful verification beside a failed build.
        failure = {
            "version": PREFIX,
            "automated_build_checks_passed": False,
            "factual_approval": False,
            "status": "Build failed; output artifacts may be partial or stale.",
            "error_type": type(exc).__name__,
            "error": str(exc),
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        }
        (ROOT / f"{PREFIX}-verification.json").write_text(
            json.dumps(failure, ensure_ascii=False, indent=2) + "\n", encoding="utf-8",
        )
        raise
    raise SystemExit(exit_code)
