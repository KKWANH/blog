"""Editable Word companion from the same CV facts/design; no Office settings changed."""
from pathlib import Path
import json, uuid, zipfile
from io import BytesIO
from lxml import etree as ET
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from results_layout import effective_spec

HERE=Path(__file__).resolve().parent
OUT=HERE.parent/'design/cv-results-2026-10-27'
ASSETS=HERE.parent/'design/cv-editorial-2026-10-27/assets'
content=json.loads((HERE/'editorial-content.json').read_text(encoding='utf-8'))
spec=effective_spec(json.loads((HERE/'editorial-design.json').read_text(encoding='utf-8')))
doc=Document()
section=doc.sections[0]
section.page_width=Mm(210);section.page_height=Mm(297)
section.top_margin=Mm(spec['margin_top_mm']);section.bottom_margin=Mm(spec['margin_bottom_mm'])
section.left_margin=section.right_margin=Mm(spec['margin_x_mm'])
section.footer_distance=Mm(12)
width=210-2*spec['margin_x_mm']
ink=spec['ink'][1:];accent=spec['accent'][1:];muted=spec['muted'][1:]
body_font=spec.get('word_body_font','Arial')
normal=doc.styles['Normal'];normal.font.name=body_font;normal.font.size=Pt(spec['body_pt'])
normal.font.color.rgb=RGBColor.from_string(ink)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=Pt(spec['body_line_pt'])
doc.core_properties.title='Kwanho Kim — Software Engineer'
doc.core_properties.author='Kwanho Kim'
doc.core_properties.subject='Experience, selected team project, skills and education'

def paragraph(parent=None,size=None,before=0,after=0,keep=False,line=None):
    parent=parent or doc
    if hasattr(parent,'paragraphs') and len(parent.paragraphs)==1 and not parent.paragraphs[0].text and parent.__class__.__name__=='_Cell':
        p=parent.paragraphs[0]
    else:p=parent.add_paragraph()
    fmt=p.paragraph_format;fmt.space_before=Pt(before);fmt.space_after=Pt(after)
    fmt.line_spacing=Pt(line or (spec['body_line_pt'] if size is None else size*1.3))
    fmt.keep_with_next=keep;fmt.widow_control=True
    return p

def run(p,text,size=None,bold=False,color=None,font=None):
    r=p.add_run(text);r.font.name=font or body_font;r.font.size=Pt(size or spec['body_pt'])
    r.bold=bold;r.font.color.rgb=RGBColor.from_string(color or ink)
    return r

def rich(p,text,phrases=(),size=None):
    cursor=0
    ranges=sorted((text.find(phrase),phrase) for phrase in phrases if phrase in text)
    for start,phrase in ranges:
        if start<cursor:continue
        run(p,text[cursor:start],size);run(p,phrase,size,bold=True);cursor=start+len(phrase)
    run(p,text[cursor:],size)

def link(p,text,url,size=9.5):
    hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');props=OxmlElement('w:rPr')
    for tag,attrs in [('w:rFonts',{'w:ascii':body_font,'w:hAnsi':body_font}),('w:sz',{'w:val':str(int(size*2))}),('w:color',{'w:val':'415748'}),('w:u',{'w:val':'single'})]:
        element=OxmlElement(tag)
        for key,value in attrs.items():element.set(qn(key),value)
        props.append(element)
    r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);hyperlink.append(r);p._p.append(hyperlink)

def border(paragraph,color='B4C3B4',side='bottom',size=4,space=4):
    props=paragraph._p.get_or_add_pPr();borders=OxmlElement('w:pBdr');edge=OxmlElement('w:'+side)
    for key,value in {'val':'single','sz':str(size),'color':color,'space':str(space)}.items():edge.set(qn('w:'+key),value)
    borders.append(edge);props.append(borders)

def table(widths,parent=None):
    if parent is None:t=doc.add_table(rows=1,cols=len(widths))
    else:t=parent.add_table(rows=1,cols=len(widths))
    t.autofit=False
    for column,mm in zip(t.columns,widths):column.width=Mm(mm)
    for cell,mm in zip(t.rows[0].cells,widths):
        cell.width=Mm(mm);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            m=OxmlElement('w:'+side);m.set(qn('w:w'),'0');m.set(qn('w:type'),'dxa');margins.append(m)
        props.append(margins)
    tr=OxmlElement('w:cantSplit');t.rows[0]._tr.get_or_add_trPr().append(tr)
    return t

def heading(text,parent=None,before=None):
    p=paragraph(parent,size=10,before=spec['section_gap_pt'] if before is None else before,after=5,keep=True,line=13)
    run(p,text.upper(),10,True,accent)
    return p

def bullet(text,phrases,parent):
    p=paragraph(parent,after=2,line=spec['body_line_pt'])
    p.paragraph_format.left_indent=Mm(3.6);p.paragraph_format.first_line_indent=Mm(-3.6)
    run(p,'• ',color='567B59');rich(p,text,phrases)
    return p

header=table([112,width-112]);left,right=header.rows[0].cells
p=paragraph(left,size=spec['name_pt'],line=spec['name_pt']*1.08)
run(p,content['name'],spec['name_pt'],True,font='Manrope')
p=paragraph(left,size=spec['role_pt'],before=5,line=18);run(p,content['role'],spec['role_pt'],False,accent)
p=paragraph(right,size=9,line=12);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run(p,'\n'.join(content['domains']),9,color=muted)
border(p,accent,'right',6,7)
p=paragraph(size=spec['contact_pt'],before=10,after=5,line=13)
p.paragraph_format.tab_stops.add_tab_stop(Mm(42))
p.paragraph_format.tab_stops.add_tab_stop(Mm(98))
p.paragraph_format.tab_stops.add_tab_stop(Mm(135))
p.paragraph_format.tab_stops.add_tab_stop(Mm(width),WD_TAB_ALIGNMENT.RIGHT)
run(p,content['location']+'\t',spec['contact_pt'],color=muted)
link(p,content['email'],'mailto:'+content['email']);run(p,'\t',spec['contact_pt'])
link(p,'kwanho.dev',content['website']);run(p,'\t',spec['contact_pt'])
link(p,'LinkedIn',content['linkedin']);run(p,'\t',spec['contact_pt']);link(p,'GitHub',content['github'])
border(p)

heading('Experience')
for index,job in enumerate(content['experience']):
    if index:paragraph(size=1,before=0,after=0,line=spec['job_gap_pt'])
    if spec['results_layout']=='index':
        t=table([spec['job_label_mm'],spec['job_column_gap_mm'],width-spec['job_label_mm']-spec['job_column_gap_mm']])
        meta,spacer,body=t.rows[0].cells
        p=paragraph(meta,size=12,line=15.5);run(p,job['company'],12,True);border(p,'89A58C','left',8,5)
        p=paragraph(meta,size=spec['meta_pt'],before=2,line=13);run(p,job['dates'],spec['meta_pt'],color=muted)
        p=paragraph(meta,size=spec['meta_pt'],line=13);run(p,job['location'],spec['meta_pt'],color=muted)
        p=paragraph(body,size=11,after=2,line=15,keep=True);run(p,job['role'],11,True)
    else:
        body=doc
        p=paragraph(keep=True,line=16)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(width),WD_TAB_ALIGNMENT.RIGHT)
        run(p,job['company']+'  ',12.5,True);run(p,job['role'],11)
        run(p,'\t'+job['dates']+' · '+job['location'],spec['meta_pt'],color=muted)
    for text in job['bullets']:bullet(text,job.get('emphasis',()),body)
    if job.get('currentFocus'):
        p=paragraph(body,before=3,line=spec['body_line_pt']);p.paragraph_format.left_indent=Mm(3.6)
        run(p,'In progress  ',bold=True,color='315D40');run(p,job['currentFocus'],color=muted)

heading('Team project & award')
for project in [p for p in content['projects'] if p.get('includeInCV',True)]:
    t=table([spec['award_label_mm'],4,width-spec['award_label_mm']-4]);flag,gap,body=t.rows[0].cells
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),accent);flag._tc.get_or_add_tcPr().append(shade)
    p=paragraph(flag,size=8,before=5,line=10);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;run(p,project['award']['credit'].upper(),8,color='F7F8ED')
    p=paragraph(flag,size=21,after=4,line=23);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;run(p,project['award']['rank'],21,True,'F7F8ED','Manrope')
    p=paragraph(body,size=11.5,line=15,keep=True);p.paragraph_format.tab_stops.add_tab_stop(Mm(width-spec['award_label_mm']-4),WD_TAB_ALIGNMENT.RIGHT)
    link(p,project['name'],project['link'],11.5);run(p,'\t'+project['dates'],spec['meta_pt'],color=muted)
    p=paragraph(body,size=10,line=14,keep=True);run(p,project['award']['scope']+' · '+project['award']['event'],10,color='31573B')
    p=paragraph(body,before=3,line=spec['body_line_pt']);run(p,project['description'])

heading('Technical skills')
t=table([width/3]*3)
for cell,row in zip(t.rows[0].cells,content['skills']):
    p=paragraph(cell,size=9.5,line=13,after=2,keep=True);run(p,row['label'],9.5,True,'365B40')
    p=paragraph(cell,size=10,line=13.5);run(p,'\n'.join(row.get('lines',[row['text']])),10)
paragraph(size=1,line=3)

t=table([(width-8)*1.35/2.35,8,(width-8)/2.35]);education,gap,training=t.rows[0].cells
heading('Education',education,before=8)
for index,row in enumerate(content['education']):
    p=paragraph(education,size=10,before=4 if index else 0,line=13,keep=True)
    run(p,row['name'],10,True)
    if row['date']:run(p,'  '+row['date'],9,color=muted)
    p=paragraph(education,size=9.5,line=12.5);run(p,'\n'.join(row.get('detailLines',[row['detail']])),9.5,color=muted)
heading('Technical training',training,before=8)
for index,row in enumerate(content['training']):
    p=paragraph(training,size=10,before=5 if index else 0,line=13,keep=True)
    run(p,row['name'],10,True);run(p,'  '+row['status'],9,color='59745A')
    p=paragraph(training,size=9.5,line=12.5);run(p,row['program']+'\n'+(row['period'] or row['location']),9.5,color=muted)
footer=paragraph(size=9.5,before=24,line=12)
run(footer,content['languages'],9.5,color=muted);border(footer,side='top',size=4,space=5)

# Embed licensed local font files in the document, rather than installing them.
for key,value in [('embedTrueTypeFonts','true'),('saveSubsetFonts','false')]:
    element=OxmlElement('w:'+key);element.set(qn('w:val'),value);doc.settings.element.append(element)
buffer=BytesIO();doc.save(buffer)
with zipfile.ZipFile(BytesIO(buffer.getvalue())) as archive:parts={name:archive.read(name) for name in archive.namelist()}
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
C='http://schemas.openxmlformats.org/package/2006/content-types'
fonts=ET.fromstring(parts['word/fontTable.xml'])
rels=ET.fromstring(parts['word/_rels/fontTable.xml.rels']) if 'word/_rels/fontTable.xml.rels' in parts else ET.Element(f'{{{P}}}Relationships')
types=ET.fromstring(parts['[Content_Types].xml'])
for index,(family,style,name) in enumerate([('Manrope','embedBold','manrope-700.ttf')],1):
    font=next((node for node in fonts if node.get(f'{{{W}}}name')==family),None)
    if font is None:font=ET.SubElement(fonts,f'{{{W}}}font',{f'{{{W}}}name':family})
    guid=uuid.uuid4();key=bytes.fromhex(guid.hex)[::-1];data=bytearray((ASSETS/name).read_bytes())
    for offset in range(32):data[offset]^=key[offset%16]
    target='fonts/'+Path(name).stem+'.odttf';rid=f'rIdCVFont{index}'
    ET.SubElement(font,f'{{{W}}}{style}',{f'{{{R}}}id':rid,f'{{{W}}}fontKey':'{'+str(guid).upper()+'}',f'{{{W}}}subsetted':'0'})
    ET.SubElement(rels,f'{{{P}}}Relationship',{'Id':rid,'Type':R+'/font','Target':target})
    ET.SubElement(types,f'{{{C}}}Override',{'PartName':'/word/'+target,'ContentType':'application/vnd.openxmlformats-officedocument.obfuscatedFont'})
    parts['word/'+target]=bytes(data)
parts['word/fontTable.xml']=ET.tostring(fonts,encoding='utf-8',xml_declaration=True)
parts['word/_rels/fontTable.xml.rels']=ET.tostring(rels,encoding='utf-8',xml_declaration=True)
parts['[Content_Types].xml']=ET.tostring(types,encoding='utf-8',xml_declaration=True)
destination=OUT/'Kwanho-Kim-CV.docx'
with zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as archive:
    for name,data in parts.items():archive.writestr(name,data)
# Structural readback confirms the file is editable Word content, not a page image.
loaded=Document(destination)
assert len(loaded.tables)>=4 and len(loaded.inline_shapes)==0
assert abs(loaded.sections[0].page_width.mm-210)<.03 and abs(loaded.sections[0].page_height.mm-297)<.03
print('Generated editable Word CV with embedded OFL fonts:',destination)
