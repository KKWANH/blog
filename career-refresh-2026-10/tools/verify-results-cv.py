from pathlib import Path
import hashlib, json, re, shutil, zipfile
from docx import Document
import pymupdf

ROOT=Path(r'C:\Projects\career-refresh-2026-10')
OUT=ROOT/'design/cv-results-2026-10-27'
REPO=Path(r'C:\Projects\blog')
content=json.loads((ROOT/'cv/editorial-content.json').read_text(encoding='utf-8'))
source=json.loads((OUT/'input-record.json').read_text(encoding='utf-8'))
input_file=Path(source['input'])
def normal(text):
    text=text.replace('\u00ad','').replace('’',"'").replace('–','-').replace('—','-')
    text=re.sub(r'\s*-\s*','-',text)
    return re.sub(r'\s+',' ',text).strip()
checks={'owner_pdf_preserved':hashlib.sha256(input_file.read_bytes()).hexdigest()==source['sha256']}
report={'checks':checks,'artifacts':{},'limits':['No physical-print or ATS ingestion test.','Actual Word rendering was checked; Google Docs import/reflow was not.','HTML/PDF uses Inter; editable Word uses the explicit Arial fallback and embedded Manrope.']}
sentences=[text for job in content['experience'] for text in job['bullets']]
sentences += [job['currentFocus'] for job in content['experience'] if job.get('currentFocus')]
sentences += [p['description'] for p in content['projects'] if p.get('includeInCV',True)]
for name,file in [('html',OUT/'linear-4.pdf'),('word',OUT/'word-proof.pdf')]:
    doc=pymupdf.open(file);text='\n'.join(page.get_text() for page in doc);flat=normal(text)
    fonts=[font for page in doc for font in page.get_fonts()]
    spans=[span for page in doc for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans'] if span['text'].strip()]
    links=[link for page in doc for link in page.get_links()]
    checks[name+'_one_a4']=len(doc)==1 and abs(doc[0].rect.width-595.276)<1 and abs(doc[0].rect.height-841.89)<1
    checks[name+'_name_first']=flat.startswith('Kwanho Kim Software Engineer')
    checks[name+'_single_contact']=text.count(content['email'])==1 and 'contact@kwanho.dev' not in text
    checks[name+'_narrative_complete']=all(normal(sentence) in flat for sentence in sentences)
    checks[name+'_job_dates_and_roles']=all(normal(job['dates']) in flat and normal(job['role']) in flat for job in content['experience'])
    checks[name+'_education_complete']=all(normal(row['name']) in flat and all(normal(line) in flat for line in row.get('detailLines',[row['detail']])) and (not row['date'] or normal(row['date']) in flat) for row in content['education'])
    checks[name+'_body_11pt']=any(abs(span['size']-11)<.15 for span in spans)
    checks[name+'_name_30pt']=any(abs(span['size']-30)<.15 for span in spans)
    checks[name+'_fonts_embedded']=bool(fonts) and all(doc.extract_font(font[0])[3] for font in fonts)
    checks[name+'_five_links']=len(links)==5 and all(link.get('uri','').startswith(('https://','mailto:')) for link in links)
    checks[name+'_gmail_link']=any(link.get('uri')=='mailto:'+content['email'] for link in links)
    checks[name+'_scope_and_training']=all(term in flat for term in ['In progress','early C++ implementation','team review','1st','digital.auto & mobility','42','SEA:ME','Completed','Figma (basic UI)'])
    checks[name+'_coursework_not_default']= 'AVL' not in text and 'C++ containers' not in text
    checks[name+'_no_unintended_font_fallback']=not any('Calibri' in font[3] for font in fonts)
    report['artifacts'][name]={'pages':len(doc),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'fonts':sorted({font[3] for font in fonts}),'links':len(links)}
    (OUT/f'final-{name}-text.txt').write_text(text,encoding='utf-8')
markdown=(ROOT/'cv/editorial-copy.md').read_text(encoding='utf-8')
checks['markdown_narrative_complete']=all(sentence in markdown for sentence in sentences)
checks['markdown_scope_complete']=all(term in markdown for term in ['In progress','1st place in the digital.auto & mobility challenge','42','SEA:ME','Completed'])
guide=(ROOT/'cv/google-docs-design-spec.md').read_text(encoding='utf-8')
checks['guide_matches_fonts']=all(term in guide for term in ['30pt','13.5pt','11pt','15.5pt','Arial'])
checks['skill_line_breaks_preserve_copy']=all(' '.join(row.get('lines',[row['text']]))==row['text'] for row in content['skills'])
word=Document(OUT/'Kwanho-Kim-CV.docx')
checks['editable_word_not_page_images']=len(word.tables)>=4 and len(word.inline_shapes)==0 and len(word.paragraphs)>10
with zipfile.ZipFile(OUT/'Kwanho-Kim-CV.docx') as archive:
    checks['word_has_no_macros']=not any('vbaProject' in name for name in archive.namelist())
    checks['word_name_font_embedded']=any(name.endswith('manrope-700.odttf') for name in archive.namelist())
report['passed']=all(checks.values())
(OUT/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
assert report['passed']
shutil.copyfile(OUT/'linear-4.pdf',OUT/'final.pdf')
shutil.copyfile(OUT/'captures/linear-4-full.png',OUT/'captures/final-full.png')
shutil.copyfile(OUT/'captures/linear-4-header.png',OUT/'captures/final-header.png')
shutil.copyfile(OUT/'final.pdf',REPO/'public/CV.pdf')
print('Verified selected PDF copied to local public/CV.pdf.')
